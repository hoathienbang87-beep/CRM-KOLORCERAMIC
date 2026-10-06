"""Emit a minimal SheetJS-compatible JSON model without modifying the XLSX."""

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def main():
    if len(sys.argv) != 2:
        raise SystemExit("usage: extract-catalog-xlsx-readonly.py <workbook.xlsx>")
    source = Path(sys.argv[1]).resolve(strict=True)
    workbook = load_workbook(source, data_only=True, read_only=False)
    output = {"SheetNames": workbook.sheetnames, "Sheets": {}}
    try:
        for worksheet in workbook.worksheets:
            sheet = {
                "!ref": f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}",
                "!merges": [
                    {
                        "s": {"r": cell_range.min_row - 1, "c": cell_range.min_col - 1},
                        "e": {"r": cell_range.max_row - 1, "c": cell_range.max_col - 1},
                    }
                    for cell_range in worksheet.merged_cells.ranges
                ],
            }
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        sheet[cell.coordinate] = {"v": json_value(cell.value)}
            output["Sheets"][worksheet.title] = sheet
    finally:
        workbook.close()
    json.dump(output, sys.stdout, ensure_ascii=False, separators=(",", ":"))


if __name__ == "__main__":
    main()
